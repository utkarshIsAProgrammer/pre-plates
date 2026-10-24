#!/usr/bin/env python3
"""Rush Bites — project presentation generator (minimal, aesthetic, 16:9)."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- palette ---
PAPER   = RGBColor(0xFA, 0xF8, 0xF4)   # warm paper background
CARD    = RGBColor(0xFF, 0xFF, 0xFF)
LINE    = RGBColor(0xE8, 0xE1, 0xD7)   # hairline borders
INK     = RGBColor(0x26, 0x23, 0x1E)   # near-black text
MUTED   = RGBColor(0x8A, 0x84, 0x7A)   # secondary text
ACCENT  = RGBColor(0xD9, 0x48, 0x0F)   # burnt orange
SOFT    = RGBColor(0xFB, 0xEA, 0xDE)   # accent tint

SERIF = "Georgia"
SANS  = "Arial"

EMU_IN = 914400
SW, SH = 13.333, 7.5          # slide size in inches
MARGIN = 0.7
CONTENT_W = SW - 2 * MARGIN   # 11.933


# ---------------------------------------------------------------- helpers ---
def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    return s


def rect(slide, x, y, w, h, fill=CARD, line=LINE, radius=0.08, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sp.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def txt(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, runs, size=12, color=INK, bold=False, font=SANS, align=PP_ALIGN.LEFT,
         italic=False, spacing=1.0, before=0, after=0):
    """runs: str, or list of (text, dict-overrides)."""
    if not tf.paragraphs[0].runs and not tf.paragraphs[0].text:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, o in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        f.name = o.get("font", font)
        f.color.rgb = o.get("color", color)
    return p


def dot_bullets(tf, items, size=11.5, gap=7, lead_color=ACCENT, text_color=INK,
                spacing=1.12, bold_lead=False):
    """Bulleted list with accent dot markers; items may be 'Lead — rest' strings."""
    first = True
    for it in items:
        if " — " in it and bold_lead:
            lead, rest = it.split(" — ", 1)
            runs = [("•  ", {"color": lead_color, "bold": True}),
                    (lead + "  —  ", {"bold": True}),
                    (rest, {"color": MUTED})]
        else:
            runs = [("•  ", {"color": lead_color, "bold": True}),
                    (it, {"color": text_color})]
        para(tf, runs, size=size, spacing=spacing,
             before=0 if first else gap)
        first = False


def header(slide, num, kicker, title):
    tf = txt(slide, MARGIN, 0.52, CONTENT_W, 0.4)
    para(tf, [(num + "  ", {"font": SERIF, "size": 16, "bold": True, "color": ACCENT}),
              (kicker.upper(), {"size": 12, "bold": True, "color": MUTED})])
    tf = txt(slide, MARGIN, 0.86, CONTENT_W, 0.75)
    para(tf, title, size=30, bold=True, font=SERIF, color=INK)
    rect(slide, MARGIN, 1.62, 0.62, 0.05, fill=ACCENT, line=None, radius=0.5)


def footer(slide, page):
    rect(slide, MARGIN, 7.06, CONTENT_W, 0.012, fill=LINE, line=None, radius=0)
    tf = txt(slide, MARGIN, 7.14, 6.0, 0.3)
    para(tf, "Rush Bites — Project Presentation", size=9, color=MUTED)
    tf = txt(slide, SW - MARGIN - 1.0, 7.14, 1.0, 0.3)
    para(tf, page, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def chip(slide, x, y, text, size=10, fill=CARD, line=LINE, color=INK, h=0.34,
         bold=True, pad=0.22):
    w = pad * 2 + 0.082 * len(text) * size / 10.0
    rect(slide, x, y, w, h, fill=fill, line=line, radius=0.5)
    tf = txt(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER)
    return x + w


def arrow(slide, x, y, w=0.3, h=0.3, glyph="→", size=14, color=MUTED):
    tf = txt(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, glyph, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)


def flow_chips(slide, x, y, labels, w=1.18, h=0.46, gap=0.34, fill=CARD,
               line=LINE, color=INK, size=10.5):
    cx = x
    for i, lab in enumerate(labels):
        if i:
            arrow(slide, cx + 0.02, y + (h - 0.3) / 2, w=gap - 0.04, color=ACCENT)
            cx += gap
        rect(slide, cx, y, w, h, fill=fill, line=line, radius=0.5)
        tf = txt(slide, cx, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, lab, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)
        cx += w
    return cx


# ---------------------------------------------------------------- deck ------
prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)

cp = prs.core_properties
cp.title = "Rush Bites — Project Presentation"
cp.author = "Rush Bites"


# ---- 1 · Title -------------------------------------------------------------
s = blank(prs)
rect(s, 9.9, -1.9, 5.6, 5.6, fill=SOFT, line=None, shape=MSO_SHAPE.OVAL)
rect(s, 12.15, 4.4, 1.9, 1.9, fill=None, line=LINE, shape=MSO_SHAPE.OVAL, line_w=1.2)
rect(s, 11.55, 5.35, 0.16, 0.16, fill=ACCENT, line=None, shape=MSO_SHAPE.OVAL)

tf = txt(s, MARGIN, 1.95, 9.5, 0.4)
para(tf, "PROJECT PRESENTATION", size=13, bold=True, color=ACCENT)
tf = txt(s, MARGIN, 2.35, 10.5, 1.5)
para(tf, [("Rush Bites", {"font": SERIF, "size": 60, "bold": True, "color": INK}),
          (".", {"font": SERIF, "size": 60, "bold": True, "color": ACCENT})])
rect(s, MARGIN, 3.75, 0.62, 0.05, fill=ACCENT, line=None, radius=0.5)
tf = txt(s, MARGIN, 4.05, 9.6, 1.0)
para(tf, "Campus dining, pre-ordered — a pre-ordering and kitchen-operations "
         "platform that replaces counter queues with scheduled, token-based pickups.",
     size=16, color=MUTED, spacing=1.25)

cx = MARGIN
for t in ["React 19", "Express 5", "MongoDB Atlas", "Clerk Auth", "Cloudinary"]:
    cx = chip(s, cx, 5.55, t, fill=CARD) + 0.16
tf = txt(s, MARGIN, 6.6, 9.0, 0.3)
para(tf, "Academic project review", size=10.5, color=MUTED)


# ---- 2 · Introduction ------------------------------------------------------
s = blank(prs)
header(s, "01", "Introduction", "Why campus dining needed a rethink")

tf = txt(s, MARGIN, 2.05, 4.5, 3.4)
para(tf, "Every lunch break, thousands descend on campus canteens within the same "
         "30–45 minute window — and the counter collapses into chaos.",
     size=15, spacing=1.3, after=10)
para(tf, [("Rush Bites", {"font": SERIF, "size": 21, "bold": True, "color": ACCENT}),
          (" moves the queue online: students order ahead, kitchens work a "
           "digital stream, and pickups happen by token number.",
           {"font": SERIF, "size": 21, "color": INK})], spacing=1.25)

tf = txt(s, MARGIN, 5.15, 4.5, 1.6)
dot_bullets(tf, [
    "One responsive web app — no installs, no POS hardware",
    "Two experiences: customer marketplace + vendor console",
], size=11.5, gap=8)

probs = [
    ("Peak-hour bottlenecks", "15–25 minutes just standing in line, inside a 30–45 min break window."),
    ("Order invisibility", "Shouted orders, lost paper slips, and zero advance demand visibility for kitchens."),
    ("Payment friction", "Cash change delays; static UPI QRs force manual amounts and disputes."),
    ("Fragmented vendors", "No central directory of open counters, live menus, prices or wait times."),
]
gx, gy, cw, ch, gap = 5.55, 2.05, 3.42, 2.28, 0.22
for i, (t, d) in enumerate(probs):
    x = gx + (i % 2) * (cw + gap)
    y = gy + (i // 2) * (ch + gap)
    rect(s, x, y, cw, ch)
    rect(s, x + 0.28, y + 0.3, 0.3, 0.05, fill=ACCENT, line=None, radius=0.5)
    tf = txt(s, x + 0.28, y + 0.5, cw - 0.56, ch - 0.7)
    para(tf, t, size=13.5, bold=True, after=5)
    para(tf, d, size=11, color=MUTED, spacing=1.18)
footer(s, "02")


# ---- 3 · Objectives --------------------------------------------------------
s = blank(prs)
header(s, "02", "Objectives", "One platform, three stakeholder pillars")

cols = [
    ("Customers", [
        "Browse menus & pre-order from classrooms, libraries or dorms",
        "Live queue position and dynamic ready-time ETA",
        "Cash on pickup or amount-locked UPI QR payments",
        "One-tap reorder with real-time stock re-validation",
        "Order-verified 1–5 star reviews for quality",
    ]),
    ("Vendors", [
        "Zero-hardware kitchen console on any phone, tablet or laptop",
        "Atomic daily token numbers (#1, #2…) that reset every morning",
        "Lifecycle queue: pending → accepted → preparing → ready → done",
        "Instant audio chimes and high-contrast new-order banners",
        "Settlement views for collected revenue vs. unpaid debt",
    ]),
    ("Institution", [
        "Decongest rush zones across campus at lunch intervals",
        "Data-driven menu planning from live sales analytics",
        "Drastically reduced food waste via demand visibility",
        "Open onboarding — any verified seller live in minutes",
    ]),
]
cw, gap = 3.84, 0.2
for i, (title, items) in enumerate(cols):
    x = MARGIN + i * (cw + gap)
    rect(s, x, 1.95, cw, 3.62)
    rect(s, x, 1.95, cw, 0.09, fill=ACCENT, line=None, radius=0)
    tf = txt(s, x + 0.28, 2.2, cw - 0.56, 0.4)
    para(tf, title, size=16, bold=True, font=SERIF)
    tf = txt(s, x + 0.28, 2.68, cw - 0.56, 2.75)
    dot_bullets(tf, items, size=10.5, gap=7)

rect(s, MARGIN, 5.85, CONTENT_W, 0.92, fill=SOFT, line=None)
tf = txt(s, MARGIN + 0.3, 5.85, CONTENT_W - 0.6, 0.92, anchor=MSO_ANCHOR.MIDDLE)
runs, first = [], True
for it in ["Role-based auth", "Multi-canteen marketplace", "Single-vendor cart",
           "Atomic daily tokens", "Amount-locked UPI QR", "Live queue polling",
           "Web Audio alerts", "Verified reviews", "Light / dark theme"]:
    if not first:
        runs.append(("   ·   ", {"color": ACCENT, "bold": True}))
    runs.append((it, {"color": INK}))
    first = False
para(tf, [("Delivered in MVP   —   ", {"bold": True, "color": ACCENT})] + runs,
     size=11.5, spacing=1.2)
footer(s, "03")


# ---- 4 · Methodology -------------------------------------------------------
s = blank(prs)
header(s, "03", "Methodology", "Decoupled architecture, disciplined order flow")

# left: architecture stack
tf = txt(s, MARGIN, 1.95, 4.35, 0.3)
para(tf, "SYSTEM ARCHITECTURE", size=11, bold=True, color=MUTED)

ax, aw = MARGIN, 4.35
rect(s, ax, 2.35, aw, 0.82)
tf = txt(s, ax + 0.25, 2.35, aw - 0.5, 0.82, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("React 19 SPA", {"bold": True, "size": 12.5}),
          ("   customers & vendors · Clerk auth · Web Audio", {"color": MUTED, "size": 10.5})])
arrow(s, ax + aw / 2 - 0.15, 3.2, glyph="↓", color=ACCENT)
rect(s, ax, 3.55, aw, 0.82)
tf = txt(s, ax + 0.25, 3.55, aw - 0.5, 0.82, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Express 5 REST API", {"bold": True, "size": 12.5}),
          ("   rate limits · JWT middleware · controllers", {"color": MUTED, "size": 10.5})])
arrow(s, ax + aw / 2 - 0.15, 4.4, glyph="↓", color=ACCENT)
rect(s, ax, 4.75, 2.08, 0.82)
tf = txt(s, ax, 4.75, 2.08, 0.82, anchor=MSO_ANCHOR.MIDDLE)
para(tf, "MongoDB Atlas", size=12, bold=True, align=PP_ALIGN.CENTER, after=2)
para(tf, "users · orders · counters", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
rect(s, ax + 2.27, 4.75, 2.08, 0.82)
tf = txt(s, ax + 2.27, 4.75, 2.08, 0.82, anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Cloudinary CDN", size=12, bold=True, align=PP_ALIGN.CENTER, after=2)
para(tf, "canteen & dish images", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
tf = txt(s, ax, 5.85, aw, 0.9)
para(tf, "Stateless REST core over HTTPS/JSON — engineered for shared campus "
         "Wi-Fi, shared NAT and flaky mobile connections.", size=10.5, color=MUTED, spacing=1.2)

# right: state machine + key solutions
rx, rw = 5.5, SW - MARGIN - 5.5
tf = txt(s, rx, 1.95, rw, 0.3)
para(tf, "ORDER LIFECYCLE STATE MACHINE", size=11, bold=True, color=MUTED)
flow_chips(s, rx, 2.32, ["pending", "accepted", "preparing", "ready", "completed"],
           w=1.18, h=0.48, gap=0.3)
tf = txt(s, rx, 2.95, rw, 0.3)
para(tf, [("cancel / reject:  ", {"bold": True, "color": MUTED}),
          ("only from pending or accepted, and only for unpaid cash orders",
           {"color": MUTED})], size=10)

tf = txt(s, rx, 3.45, rw, 0.3)
para(tf, "KEY ENGINEERING SOLUTIONS", size=11, bold=True, color=MUTED)
tf = txt(s, rx, 3.82, rw, 2.9)
dot_bullets(tf, [
    "Atomic daily tokens — Mongo $inc on {canteen, day}: race-free, zero-gap #1, #2… with no distributed locks.",
    "Amount-locked UPI deep links — NPCI intent URI rendered as client-side SVG QR; totals pre-filled, no manual typing.",
    "Live queue ETA — prep time + 2 min × orders-ahead, computed from one source of truth for both screens.",
    "Campus-NAT rate limiting — per-user buckets with IP-subnet fallback and tiered thresholds (orders, auth, uploads).",
    "Offline resilience — 15 s health heartbeat, connection banners, Cloudinary→Base64 image fallback.",
], size=10.5, gap=8, bold_lead=True, spacing=1.12)
footer(s, "04")


# ---- 5 · Technologies ------------------------------------------------------
s = blank(prs)
header(s, "04", "Technologies Used", "A deliberate, modern stack")

groups = [
    ("Frontend", ["React 19", "Vite 7", "Clerk React", "react-qr-code",
                  "Web Audio API", "Vanilla CSS3"]),
    ("Backend", ["Node.js 20", "Express 5", "Clerk Express", "express-rate-limit",
                 "CORS", "dotenv"]),
    ("Data & Media", ["MongoDB Atlas 7", "Mongoose 9", "Aggregation pipelines",
                      "Cloudinary CDN", "Compound indexes"]),
    ("Design & UX", ["CSS variable theming", "Dark / light modes",
                     "Print stylesheet for token slips", "Responsive layout"]),
]
cw, ch, gap = 5.86, 1.92, 0.21
for i, (title, items) in enumerate(groups):
    x = MARGIN + (i % 2) * (cw + gap)
    y = 1.95 + (i // 2) * (ch + gap)
    rect(s, x, y, cw, ch)
    tf = txt(s, x + 0.28, y + 0.22, cw - 0.56, 0.35)
    para(tf, title, size=14.5, bold=True, font=SERIF)
    cx, cy = x + 0.28, y + 0.68
    for t in items:
        w = 0.44 + 0.082 * len(t)
        if cx + w > x + cw - 0.28:
            cx = x + 0.28
            cy += 0.46
        chip(s, cx, cy, t, size=10, h=0.36, fill=PAPER, line=LINE)
        cx += w + 0.14

rect(s, MARGIN, 6.0, CONTENT_W, 0.78, fill=SOFT, line=None)
tf = txt(s, MARGIN + 0.3, 6.0, CONTENT_W - 0.6, 0.78, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Why this stack — ", {"bold": True, "color": ACCENT}),
          ("zero-dependency audio synthesis (no mp3 404s), atomic Mongo operations for "
           "concurrency, CDN images auto-capped at 1200px, and sub-second loads on 4G.",
           {"color": INK})], size=11.5, spacing=1.15)
footer(s, "05")


# ---- 6 · Implementation ------------------------------------------------------
s = blank(prs)
header(s, "05", "Implementation", "From schema to screen")

cards = [
    ("Data layer", [
        "6 Mongoose models — User, Canteen, Menu, Order, Counter, Review",
        "Compound unique index {canteen, day} powers daily tokens",
        "Denormalized rating aggregates on canteens",
        "Structured payment sub-document with UTR reference",
    ]),
    ("API layer", [
        "REST modules — /auth, /canteens, /menu, /orders, /reviews, /uploads",
        "Role-scoped vendor routes for queue, stats and status",
        "Route-order guard: /mine declared before /:id",
        "Server-enforced single-canteen cart invariant",
    ]),
    ("Client layer", [
        "Context chain — Profile → Canteen → Cart, no Redux",
        "Role-adaptive tabs for customers vs. vendors",
        "PayModal builds NPCI URI and renders SVG QR",
        "tone() engine synthesizes chimes; @media print token slips",
    ]),
]
cw, gap = 3.84, 0.2
for i, (title, items) in enumerate(cards):
    x = MARGIN + i * (cw + gap)
    rect(s, x, 1.95, cw, 3.28)
    rect(s, x, 1.95, cw, 0.09, fill=ACCENT, line=None, radius=0)
    tf = txt(s, x + 0.28, 2.2, cw - 0.56, 0.4)
    para(tf, title, size=16, bold=True, font=SERIF)
    tf = txt(s, x + 0.28, 2.68, cw - 0.56, 2.4)
    dot_bullets(tf, items, size=10.5, gap=8)

tf = txt(s, MARGIN, 5.5, CONTENT_W, 0.3)
para(tf, "THE VERIFIED USER JOURNEY", size=11, bold=True, color=MUTED)
steps = ["Order placed", "Kitchen chime", "Accept · prepare", "Ready chime",
         "Token #14 pickup", "Payment confirmed", "5★ review"]
flow_chips(s, MARGIN, 5.85, steps, w=1.5, h=0.46, gap=0.22, fill=CARD, size=9.5)
footer(s, "06")


# ---- 7 · Outcomes ------------------------------------------------------------
s = blank(prs)
header(s, "06", "Outcomes", "Measured impact on campus dining")

metrics = [
    ("~100%", "queue wait eliminated — walk straight to pickup"),
    ("3×", "kitchen throughput — 80–120 orders/hour vs. 25–35"),
    ("0%", "order errors — exact items, quantities and notes"),
    ("₹0", "hardware spend — runs on any existing device"),
]
cw, gap = 2.83, 0.2
for i, (big, lab) in enumerate(metrics):
    x = MARGIN + i * (cw + gap)
    rect(s, x, 1.95, cw, 1.72)
    tf = txt(s, x + 0.24, 2.12, cw - 0.48, 0.85)
    para(tf, big, size=36, bold=True, font=SERIF, color=ACCENT)
    tf = txt(s, x + 0.24, 3.0, cw - 0.48, 0.6)
    para(tf, lab, size=10, color=MUTED, spacing=1.15)

rect(s, MARGIN, 3.95, 5.86, 2.75)
tf = txt(s, MARGIN + 0.28, 4.2, 5.3, 0.35)
para(tf, "Built for real campuses", size=15, bold=True, font=SERIF)
tf = txt(s, MARGIN + 0.28, 4.68, 5.3, 1.9)
dot_bullets(tf, [
    "Zero-hardware parity — phone as order receiver, tablet as 'Now Serving' display",
    "Dual-mode settlement — locked UPI totals plus tracked unpaid cash debt",
    "Survives shared-NAT Wi-Fi, silent JWT rotation and database drops",
], size=10.5, gap=9)

rect(s, 6.76, 3.95, 5.86, 2.75)
tf = txt(s, 7.04, 4.2, 5.3, 0.35)
para(tf, "Roadmap", size=15, bold=True, font=SERIF)
tf = txt(s, 7.04, 4.68, 5.3, 1.9)
dot_bullets(tf, [
    "WebSockets / SSE — sub-second push updates for kitchens",
    "Web Push + WhatsApp alerts — order-ready pings beyond the browser",
    "RFID smartcard wallets — direct campus meal-plan deduction",
    "Predictive prep analytics — forecast demand before the bell rings",
], size=10.5, gap=9)
footer(s, "07")


# ---- 8 · Thank you -----------------------------------------------------------
s = blank(prs)
rect(s, -1.7, 4.6, 4.6, 4.6, fill=SOFT, line=None, shape=MSO_SHAPE.OVAL)
tf = txt(s, 1.7, 2.5, 9.93, 1.2)
para(tf, [("Thank you", {"font": SERIF, "size": 54, "bold": True, "color": INK}),
          (".", {"font": SERIF, "size": 54, "bold": True, "color": ACCENT})],
     align=PP_ALIGN.CENTER)
rect(s, SW / 2 - 0.31, 3.85, 0.62, 0.05, fill=ACCENT, line=None, radius=0.5)
tf = txt(s, 1.7, 4.1, 9.93, 0.5)
para(tf, "Rush Bites — campus dining, pre-ordered", size=15, color=MUTED,
     align=PP_ALIGN.CENTER)
tf = txt(s, 1.7, 4.7, 9.93, 0.4)
para(tf, "Questions & discussion welcome", size=11, bold=True, color=ACCENT,
     align=PP_ALIGN.CENTER)

labels = ["Introduction", "Objectives", "Methodology", "Technologies",
          "Implementation", "Outcomes"]
cx = 2.45
for t in labels:
    cx = chip(s, cx, 5.5, t, size=9.5, h=0.32, fill=CARD) + 0.14

prs.save("RushBites_Project_Presentation.pptx")
print("Saved RushBites_Project_Presentation.pptx with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
